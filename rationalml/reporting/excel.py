"""Write result tables to Excel without accessing training data or the model."""

import json
import os
from datetime import date, datetime, time
from decimal import Decimal
from math import isfinite
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import TYPE_CHECKING, Any

import numpy as np
import pandas as pd

from ..exceptions import ConfigurationError, DataValidationError, MissingDependencyError
from ..tasks import TaskType
from .config import ExcelReportConfig

if TYPE_CHECKING:
    from ..result import AutoMLResult

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
_PERCENTAGES = {"positive_rate", "population_share", "positive_capture", "cumulative_positive_capture"}
_INTEGERS = {"rank", "segment", "bin", "count", "positives", "n_trials_completed"}


def _json_default(value: Any) -> Any:
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, (set, frozenset)):
        return sorted(value, key=str)
    if isinstance(value, (Path, date, time, Decimal)):
        return str(value)
    return f"<{type(value).__module__}.{type(value).__qualname__}>"


def _excel_value(value: Any) -> Any:
    if isinstance(value, np.generic):
        value = value.item()
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, float) and not isfinite(value):
        if np.isnan(value):
            return None
        raise DataValidationError("Excel cannot represent infinite numeric values.")
    if isinstance(value, (datetime, time)) and value.tzinfo is not None:
        return value.isoformat()
    if isinstance(value, (str, bool, int, float, Decimal, datetime, date, time)):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=_json_default)


def _decimal_format(places: int) -> str:
    return "0" + ("." + "0" * places if places else "")


def _number_format(column: str, value: Any, config: ExcelReportConfig) -> str:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return "General"
    if column in _PERCENTAGES:
        return _decimal_format(config.percentage_places) + "%"
    if column in _INTEGERS or isinstance(value, int):
        return "0"
    if column in {"lift", "cumulative_lift"}:
        return "0.00"
    return _decimal_format(config.decimal_places)


def _check_size(name: str, table: pd.DataFrame, include_index: bool) -> None:
    if len(table) + 1 > EXCEL_MAX_ROWS:
        if name == "Predictions":
            raise DataValidationError("Predictions exceed Excel row limit; disable include_predictions or export separately.")
        raise DataValidationError(f"Sheet {name!r} exceeds Excel row limit ({EXCEL_MAX_ROWS}, including header).")
    if len(table.columns) + int(include_index) > EXCEL_MAX_COLUMNS:
        raise DataValidationError(f"Sheet {name!r} exceeds Excel column limit ({EXCEL_MAX_COLUMNS}).")


def _write_table(workbook: Any, name: str, table: pd.DataFrame, config: ExcelReportConfig, include_index: bool) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(name)
    headers = ([table.index.name or "index"] if include_index else []) + list(table.columns)
    sheet.append([str(header) for header in headers])
    for cell in sheet[1]:
        cell.data_type = "s"
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EDF2")
    widths = [len(str(header)) for header in headers]
    for row_number, row in enumerate(table.itertuples(index=include_index, name=None), start=2):
        for position, (header, raw_value) in enumerate(zip(headers, row), start=1):
            value = _excel_value(raw_value)
            if isinstance(value, str) and len(value) > 32_767:
                raise DataValidationError(f"Sheet {name!r} contains text exceeding Excel's 32767-character cell limit.")
            cell = sheet.cell(row=row_number, column=position, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # Preserve labels literally, including strings starting with '='.
            cell.number_format = _number_format(str(header), value, config)
            widths[position - 1] = max(widths[position - 1], min(48, len(str(value)) if value is not None else 0))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for position, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = min(50, max(12, width + 2))


def _summary(result: "AutoMLResult", class_label: object | None) -> pd.DataFrame:
    from .. import __version__

    best = result.leaderboard.loc[result.best_model_name]
    items = {
        "RationalML version": __version__, "Task": result.task.value, "Target": result.target,
        "Primary metric": result.primary_metric, "Best model": result.best_model_name,
        "Baseline name": result.baseline_name, "Baseline score": result.baseline_score,
        "Best CV score": best["cv_score"], "CV std": best["cv_std"],
        "Improvement vs baseline": best["improvement_vs_baseline"],
        "Test primary metric": result.test_metrics.get(result.primary_metric),
        "Train rows": len(result.train_indices), "Test rows": len(result.test_indices),
        "Number of features": len(result.feature_names), "CV folds": result.config.cv,
        "Optuna trials requested": result.config.n_trials, "Random state": result.config.random_state,
    }
    if result.task is TaskType.BINARY:
        items.update({"Positive class": result.positive_class, "Negative class": result.negative_class})
    elif result.task is TaskType.MULTICLASS:
        items.update({"Number of classes": len(result.classes_), "Classes": result.classes_.tolist()})
        if class_label is not None:
            items["Ranking class"] = class_label
    if result.task is TaskType.BINARY or (result.task is TaskType.MULTICLASS and class_label is not None):
        summary = result.calibration_summary(n_bins=10, class_label=class_label)
        items.update({"Brier score": summary["brier_score"],
                      "Expected calibration error": summary["expected_calibration_error"],
                      "Max calibration error": summary["max_calibration_error"]})
    return pd.DataFrame(items.items(), columns=["Item", "Value"])


def export_excel(
    result: "AutoMLResult", path: str | Path, *, config: ExcelReportConfig | None = None,
    class_label: object | None = None,
) -> Path:
    """Serialize stored results; no parent creation, only .xlsx, overwrite on success."""
    try:
        from openpyxl import Workbook
    except ImportError as error:
        raise MissingDependencyError(
            'Excel export requires openpyxl. Install RationalML[excel]: pip install "rationalml[excel]".'
        ) from error
    if config is None:
        config = ExcelReportConfig()
    elif not isinstance(config, ExcelReportConfig):
        raise ConfigurationError("config must be an ExcelReportConfig or None.")
    if not isinstance(path, (str, Path)) or Path(path).suffix.lower() != ".xlsx":
        raise ConfigurationError("Excel export requires a path with the .xlsx extension.")
    path = Path(path)
    if not path.parent.is_dir():
        raise FileNotFoundError(f"Excel output parent directory does not exist: {path.parent}")
    if path.is_dir():
        raise IsADirectoryError(f"Excel output path is a directory: {path}")
    if result.task is not TaskType.MULTICLASS and class_label is not None:
        raise ConfigurationError("class_label is only used for multiclass ranking.")
    importance = result.feature_importance
    if importance.empty and len(importance.columns) == 0:
        importance = pd.DataFrame({"message": ["Feature importance unavailable."]})
    tables = [
        ("Summary", _summary(result, class_label), False),
        ("Leaderboard", result.leaderboard.rename_axis("model").reset_index(), False),
        ("Metrics", pd.DataFrame([(name, value, name == result.primary_metric) for name, value in result.test_metrics.items()],
                                 columns=["metric", "value", "primary"]), False),
        ("Feature Importance", importance, False),
        ("Hyperparameters", pd.DataFrame(result.best_params.items(), columns=["parameter", "value"]), False),
    ]
    if result.task is not TaskType.MULTICLASS or class_label is not None:
        tables.extend([
            ("Ranking", result.ranking_table(class_label=class_label), False),
            ("Top Segment", result.top_segment(config.top_fraction, class_label=class_label), True),
        ])
    if result.task is TaskType.BINARY or (result.task is TaskType.MULTICLASS and class_label is not None):
        tables.append(("Calibration", result.calibration_table(n_bins=10, class_label=class_label), False))
    if config.include_predictions:
        tables.append(("Predictions", result.test_predictions, True))
    for name, table, include_index in tables:
        _check_size(name, table, include_index)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, table, include_index in tables:
        _write_table(workbook, name, table, config, include_index)
    # Close the temporary handle before openpyxl opens it (required on Windows).
    with NamedTemporaryFile(dir=path.parent, prefix=f".{path.name}.", suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)
        workbook.close()
    return path
